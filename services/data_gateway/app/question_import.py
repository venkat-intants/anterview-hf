"""One spreadsheet reader for every question-import path.

Two surfaces import questions from a file an HR manager built in Excel: an
EXAM's section (``routers/hr_exams.py``, since PH3) and now a QUESTION BANK
(``routers/question_banks.py``). The reading, the "Correct" resolution and the
per-row error reporting are identical problems, and this module is the single
copy of them — a second reader would drift, and the way it would drift is
silent: a file that imports into an exam but is refused by a bank, or accepted
with a different correct answer, with nothing to say which reader was right.

WHAT THE TWO LAYOUTS SHARE, AND WHERE THEY DIFFER

Both start with the exam template's seven columns, in that order:

    Question | Option A | Option B | Option C | Option D | Correct | Points

A bank row may carry three more, because a bank question is reusable and so
needs the facts that make it findable and gradeable later:

    ... | Difficulty | Language | Competency

The extra columns are OPTIONAL in the file. A bank import with the plain
seven-column exam template works and takes the caller's defaults, so an HR
manager who already has exam sheets can load them straight into a bank — which
is most of the point.

"Correct" names the right option by letter (A-D), by number (1-4), or by
repeating the option's own text. All three because all three are what people
actually type, and refusing two of them would send HR back to re-key a file
that is not wrong, only differently spelled.

NO VALIDATION LIVES HERE THAT THE CALLER CANNOT SEE. This module returns plain
dicts and a per-row error list; it raises nothing, writes nothing and knows
nothing about either destination's rules (an exam's section lock, a bank's
review lifecycle). The caller validates its own domain, which is why a bank can
land these as drafts needing review while an exam appends them directly.
"""

from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from typing import Any

import structlog

#: The exam template's columns, and the first seven of a bank's.
TEMPLATE_HEADER: tuple[str, ...] = (
    "Question", "Option A", "Option B", "Option C", "Option D", "Correct", "Points",
)

#: A bank row's extra columns. Optional — a seven-column exam sheet imports fine.
BANK_EXTRA_HEADER: tuple[str, ...] = ("Difficulty", "Language", "Competency")

BANK_TEMPLATE_HEADER: tuple[str, ...] = TEMPLATE_HEADER + BANK_EXTRA_HEADER

#: Upload cap. A DoS guard, not a product limit: 2 MB of .xlsx is many
#: thousands of rows, and MAX_ROWS below is the limit a person actually meets.
MAX_IMPORT_BYTES: int = 2_000_000

#: Rows parsed from one file. Bounds the per-row work a single request can buy
#: — each imported bank question is an ORM insert plus an event row, so an
#: unbounded file is a slow transaction holding locks, not just a big response.
MAX_ROWS: int = 1_000

#: Cells kept per row. A row wider than the template carries nothing we read,
#: and a crafted ``<dimension ref="A1:XFD1048576"/>` makes openpyxl's read-only
#: reader pad EVERY row out to 16,384 cells — which is how a 25 KB upload
#: allocated 2.7 GB. Slicing here makes the padding free.
MAX_COLS: int = len(BANK_TEMPLATE_HEADER)

#: Rows the reader will LOOK AT, as opposed to the MAX_ROWS it will use. Excel
#: emits trailing blank rows for any touched cell, so the gap between the two is
#: deliberate slack — but it has to be finite, or a sheet declaring a million
#: empty rows costs a million iterations before the first question is found.
MAX_SCAN_ROWS: int = MAX_ROWS * 5

DIFFICULTIES: frozenset[str] = frozenset({"easy", "medium", "hard"})
LANGUAGES: frozenset[str] = frozenset({"en", "hi", "te"})


class SpreadsheetError(Exception):
    """The file could not be read at all. Carries a sentence for a person."""


def competency_slug(name: str) -> str:
    """An id for a competency name, always matching ``^[a-z0-9_]{1,80}$``.

    ``question_banks.validate_competencies`` enforces that pattern, and this
    function is what guarantees a row can satisfy it.

    THE BUG IT FIXES (review 2026-10-05). The slug was built with
    ``ch.isalnum()``, which is True for Devanagari and Telugu letters — so a
    Hindi or Telugu competency name produced a non-ASCII id, validation raised
    INSIDE the import loop, and the whole file rolled back naming no row at all.
    EN/HI/TE are the Day-1 languages (CLAUDE.md hard constraint 5), so the
    languages this platform exists to serve were the ones that broke it.

    A name with no usable ASCII keeps its HUMAN NAME intact — that is what HR
    and the picker display — and takes a stable hash as its id. Opaque, but
    valid, deterministic (the same name always gives the same id, so two rows
    tagged alike stay alike), and honest: inventing a transliteration would put
    a guess in the data.
    """
    ascii_slug = "".join(
        ch if ("a" <= ch <= "z" or "0" <= ch <= "9") else "_" for ch in name.casefold()
    ).strip("_")[:80]
    # "_" * n strips to empty, so this also catches a name of only separators.
    if ascii_slug.strip("_"):
        return ascii_slug
    digest = hashlib.sha256(name.strip().casefold().encode("utf-8")).hexdigest()[:12]
    return f"comp_{digest}"


@dataclass(frozen=True)
class RowError:
    """One row we could not use, and why — in the HR manager's terms."""

    row: int
    message: str

    def as_dict(self) -> dict[str, Any]:
        return {"row": self.row, "message": self.message}


def correct_to_index(
    raw: str, options: list[str], *, columns: list[str] | None = None
) -> int:
    """Resolve the 'Correct' cell to a 0-based index into ``options``.

    Accepts 'A'-'D', '1'-'4', or the option's own text (case-insensitive).

    ``columns`` is the row's four Option cells INCLUDING any blanks, and passing
    it is what makes a letter mean a COLUMN. ``options`` is the blanks-removed
    list that gets stored.

    THE BUG THIS SIGNATURE EXISTS TO FIX (found in review 2026-10-05, present in
    the exam importer since PH3). A letter used to index ``options`` directly,
    so a row with a blank middle column silently marked the WRONG answer:

        Which wire is live? | red | <blank> | blue | green | C

    compacts to ``[red, blue, green]``, and 'C' -> index 2 -> "green". The sheet
    says column C, which is "blue". Nothing was reported — a wrong graded fact
    imported in silence, into a library that decides who gets interviewed. A
    letter now resolves against the columns and is then mapped onto the stored
    list, and naming an EMPTY column is an error rather than a shift.

    Without ``columns`` the old positional behaviour is kept, so a caller that
    has no column context (there is none today) is unchanged.
    """
    s = (raw or "").strip()
    if not s:
        raise ValueError("missing correct answer")

    def _by_position(idx: int, shown: str) -> int:
        """Map a 0-based COLUMN index onto the stored, compacted options."""
        if columns is None:
            if 0 <= idx < len(options):
                return idx
            raise ValueError(f"correct '{shown}' is out of range for {len(options)} options")
        if not 0 <= idx < len(columns):
            raise ValueError(f"correct '{shown}' is out of range for {len(columns)} columns")
        cell = (columns[idx] or "").strip()
        if not cell:
            letter = chr(ord("A") + idx)
            raise ValueError(f"correct '{shown}' names option {letter}, which is empty")
        return options.index(cell)

    if len(s) == 1 and s.isalpha():  # 'A'..'D'
        return _by_position(ord(s.upper()) - ord("A"), s)
    if s.isdigit():  # '1'..'4'
        return _by_position(int(s) - 1, s)
    for i, o in enumerate(options):  # literal option text
        if o.strip().casefold() == s.casefold():
            return i
    raise ValueError(f"correct answer '{s}' matches no option")


log = structlog.get_logger(__name__)

#: The one sentence every unreadable-file path answers with. Written once because
#: it was written three times, and two of those were a copy of the first.
_UNREADABLE = "Could not read that file — is it a valid .xlsx or .csv?"

#: An OOXML package starts with the local-file-header magic of a zip, AT OFFSET 0.
#: Offset 0 matters: ``zipfile`` finds the end-of-central-directory record by
#: scanning backwards, so an archive appended to another file is still a readable
#: archive. Without this, a body beginning ``%PDF-`` with an xlsx appended answers
#: True to "is this a workbook?" AND to "is this a PDF?", and which parser runs
#: depends on which sniffer is asked first.
_ZIP_MAGIC = b"PK\x03\x04"

#: The entry that makes a zip a WORKBOOK rather than some other OOXML document.
#: Structural, like ``document_storage.sniff_corpus``'s ``word/document.xml``
#: check: a .docx is also a zip, and handing one to ``load_workbook`` is not
#: meaningfully better than handing it a CSV.
_XLSX_REQUIRED_ENTRY = "xl/workbook.xml"

#: Format headers that are definitely not a spreadsheet or a CSV.
#:
#: WHY A DENY-LIST AND NOT "DOES THIS LOOK LIKE TEXT". Content dispatch sends every
#: non-workbook upload to ``csv.reader``, and Python 3.12's csv module no longer
#: rejects NUL bytes — so a PDF posted to the importer came back as a hundred rows
#: of mojibake and a 200 with a hundred row errors, where before it was one clean
#: 400 (review 2026-10-06). The obvious guards are both wrong, measured:
#:
#:   * A blanket "any NUL" check breaks a property that exists on purpose —
#:     ``parse_mcq_rows`` names the ROW a NUL appears in and imports the rest,
#:     because Postgres refuses a NUL in a text column and left in it surfaced as a
#:     500 at flush time, failing a 200-row import over one dirty cell.
#:   * A NUL- or control-byte DENSITY check cannot see a PDF at all: measured on a
#:     real one, 0.00% NUL and 0.00% control bytes in the first 8 KB, because a PDF
#:     is mostly ASCII. Meanwhile a 93-byte CSV with one stray NUL scores 1.08%,
#:     above any threshold that would catch a PNG.
#:
#: So the discriminating signal is the header, which is what
#: ``document_storage.sniff`` already uses — this is that idea in reverse, naming
#: what is definitely not a spreadsheet.
#:
#: Narrower than a blanket check on purpose: an UNRECOGNISED binary format still
#: reaches the CSV branch and produces per-row errors, which is the pre-existing
#: behaviour for any unparseable content and is better than losing the row-level
#: NUL message.
#: THE RULE FOR THIS LIST, and it is not negotiable by adding "one more obvious
#: header": a magic qualifies only if it contains a byte that CANNOT APPEAR IN TEXT.
#:
#: The first version of this list ignored that and included ``BM`` (BMP), ``MZ``
#: (exe), ``GIF8``, ``BZh`` and ``Rar!``, all printable ASCII. Measured against it,
#: these ordinary question sheets were refused as binary:
#:
#:     BMI,Option A,Option B,...            -> 400
#:     MZN is the currency of?,Mozambique   -> 400
#:     BZh is the magic of?,bzip2,gzip      -> 400
#:
#: ``_skip_header`` supports header-less sheets on purpose, and "BMI" is a plausible
#: first question for the healthcare family in ``shared/intelligence/``. The user
#: got a sentence about .xlsx validity, which says nothing about the content — an
#: outage they could not diagnose or report (security review 2026-10-06).
#:
#: ``%PDF-`` is the ONE exception to the rule. It is printable, and it is kept
#: because a resume posted to the question importer is the actual reported harm,
#: five characters is long enough to be distinctive, and no question begins "%PDF-".
_BINARY_MAGICS: tuple[bytes, ...] = (
    b"%PDF-",                # PDF — the stated exception; see the rule above
    b"\x89PNG\r\n\x1a\n",    # PNG
    b"\xff\xd8\xff",         # JPEG
    b"\xd0\xcf\x11\xe0",     # OLE2: .xls, .doc — openpyxl never read these
    b"\x7fELF",              # Linux executable
    b"\x1f\x8b",             # gzip
    b"\xfd7zXZ",             # xz
    b"7z\xbc\xaf",           # 7-zip
    b"\xff\xfe",             # UTF-16 LE BOM — text, but csv reads it as NULs
    b"\xfe\xff",             # UTF-16 BE BOM
)

#: Dropped deliberately, each because it is printable ASCII and collided with a real
#: question: BM (BMP) with "BMI", MZ (exe) with "MZN", GIF8/BZh/Rar! with questions
#: that name those formats. A BMP or an executable posted to a question importer now
#: reaches the CSV branch and is reported per row, which is the pre-existing
#: behaviour for any unparseable content and is the price of not refusing "BMI".
_DELIBERATELY_NOT_MAGICS: tuple[bytes, ...] = (b"BM", b"MZ", b"GIF8", b"BZh", b"Rar!")


def _looks_binary(content: bytes) -> bool:
    """True when *content* announces itself as a format that is not a spreadsheet.

    A ZIP is included implicitly: a zip that is a workbook never reaches here (the
    caller checks :func:`looks_like_xlsx` first), and a zip that is NOT a workbook —
    a .docx, say — is an archive, so reading it as comma-separated text produces
    nothing but noise.
    """
    return content.startswith(_BINARY_MAGICS) or content.startswith(_ZIP_MAGIC)


def looks_like_xlsx(content: bytes) -> bool:
    """True when *content* really is an OOXML workbook, by its bytes.

    Named for what it does: it checks the zip magic and one entry, not validity.
    ``load_workbook`` is still the thing that decides whether the workbook can be
    read.

    Asked instead of "is this text?" on purpose. The CSV branch decodes with
    ``errors="replace"``, so no encoding fails there: a Windows-1252 export from a
    non-English Excel imports with one replacement character. A strict utf-8 sniff
    — the shape ``sniff_corpus`` uses — would start REFUSING those files, trading
    an attack control for an HR manager who can no longer upload.

    ``.xls`` (OLE2, not a zip) is False here and was already refused by
    ``load_workbook`` with the same message, so that path is unchanged.

    DELIBERATELY STRICTER THAN openpyxl, which is worth knowing before widening it.
    openpyxl does not look for the literal name ``xl/workbook.xml``: it reads the
    main part's path out of ``[Content_Types].xml``, so an OOXML-legal workbook may
    declare it anywhere, and this predicate refuses one that does (security review
    2026-10-06). Accepted, not fixed, because Excel, LibreOffice and Google Sheets
    all write ``xl/workbook.xml`` — so the cost falls only on programmatic
    generators — and the alternative is parsing the manifest in order to decide
    which parser to use, which puts attacker-controlled XML upstream of the
    dispatch decision. Do NOT widen this to "any zip": keeping a .docx out of
    ``load_workbook`` is what the required entry is for.
    """
    if not content.startswith(_ZIP_MAGIC):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            return _XLSX_REQUIRED_ENTRY in z.namelist()
    except (zipfile.BadZipFile, OSError):
        # A truncated or lying zip is not a workbook. ``OSError`` is kept beside
        # ``BadZipFile`` — ``sniff_corpus`` catches only the latter — because a
        # PREDICATE that raises is worse than one that over-catches: this is
        # called before anything has decided what the file is, and an exception
        # here would escape as a 500.
        return False


def _claims_to_be_a_workbook(filename: str | None) -> bool:
    """True when the NAME says Excel. Used only to word an error, never to dispatch.

    ``.xls`` is included although openpyxl has never been able to read one: an
    operator who uploads a real ``.xls`` deserves the sentence naming the two
    formats that do work, which is what they got before this module looked at
    content at all.
    """
    return (filename or "").lower().endswith((".xlsx", ".xlsm", ".xls"))


def assert_xlsx_parser_hardened() -> None:
    """Raise unless openpyxl's XML hardening is active. Cheap; safe to call often.

    ``.xlsx`` is a zip of XML. openpyxl defuses its parser only when ``defusedxml``
    imports AND ``OPENPYXL_DEFUSEDXML`` is unset or exactly ``"True"``, and it
    computes that once at import into ``openpyxl.xml.DEFUSEDXML``. With it False
    the ``read_only=True`` reader below runs on the stdlib expat parser, which
    expands internal entities — measured on this fixture: a laced
    ``<!ENTITY>`` chain returns 10,000 characters from a one-kilobyte upload with
    the flag off, and is refused with it on.

    WHY THE FLAG AND NOT ``import defusedxml``. The import check misses the second
    failure mode entirely: ``OPENPYXL_DEFUSEDXML=true`` — lowercase, the spelling
    an operator would guess — disables the hardening with the package installed and
    correct, because openpyxl compares against the exact string ``"True"``.

    Called from TWO places on purpose: ``app/main.py``'s lifespan, so a
    misconfigured deploy fails at boot rather than at the first upload; and
    :func:`_read_workbook`, so the guarantee holds at the point of use for anything
    that never ran a lifespan (a script, a worker, a test).

    This lived on ``Settings`` as a model validator and moved here (review
    2026-10-06): it read nothing off ``self``, and importing openpyxl at
    ``app.config`` import time cost ~180 ms to EVERY importer of ``settings`` —
    including ``alembic/env.py``, so a database migration would have refused to run
    over a spreadsheet parser's environment variable.
    """
    try:
        import openpyxl.xml  # noqa: PLC0415 — probed, not required; see below
    except ImportError:
        # openpyxl missing is a different failure, reported per request by
        # _read_workbook ("Excel support is not installed on the server"). Not this
        # check's business, and raising here would turn it into a boot failure.
        #
        # Note what this branch is and is not: `requirements.txt` PINS
        # openpyxl==3.1.5, so it is unreachable in the built image. It exists for a
        # script or a slimmed environment, not as a "supported degraded
        # deployment" — an earlier version of this docstring called it that, which
        # contradicted the pin (security review 2026-10-06).
        return
    if not openpyxl.xml.DEFUSEDXML:
        raise RuntimeError(
            "openpyxl is installed but its XML hardening is OFF "
            "(openpyxl.xml.DEFUSEDXML is False), so uploaded .xlsx files would be "
            "parsed with entity expansion enabled — a kilobyte of billion-laughs "
            "becomes gigabytes of heap. Either install defusedxml (pinned in "
            "requirements.txt) or unset OPENPYXL_DEFUSEDXML, which disables the "
            "hardening unless it is exactly the string 'True'."
        )


def _read_csv(content: bytes) -> list[list[str]]:
    """Rows from a CSV body. Everything that is not a workbook comes here."""
    if _looks_binary(content):
        # See _BINARY_MAGICS: without this, content dispatch turned a PDF posted to
        # the importer into a hundred rows of mojibake and a 200. A stray NUL in an
        # otherwise-textual CSV is NOT caught here — that stays a per-row message,
        # which is a property parse_mcq_rows has on purpose.
        raise SpreadsheetError(_UNREADABLE)
    text = content.decode("utf-8-sig", errors="replace")
    try:
        return _bounded([(c or "") for c in row] for row in csv.reader(io.StringIO(text)))
    except csv.Error as exc:
        # csv raises on a field over its 131072-char limit, and on a few
        # malformed-quoting shapes. Uncaught it was a 500 on a file the uploader
        # could have fixed.
        #
        # The reason used to be interpolated verbatim, which handed an HR manager
        # Python's own words — "new-line character seen in unquoted field - do you
        # need to open the file with newline=''?" is a library API hint, not advice
        # anybody can act on (security review 2026-10-06). The sentence they get is
        # about their file; the library's text goes to the log.
        log.warning("question_import.csv_unreadable", error=str(exc))
        raise SpreadsheetError(
            "Could not read that CSV — check for an unclosed quote or a cell "
            "larger than 128 KB."
        ) from exc


def _read_workbook(content: bytes) -> list[list[str]]:
    """Rows from an .xlsx body, bounded as it streams."""
    try:
        import openpyxl  # noqa: PLC0415 — lazy: only the Excel path needs it
    except ImportError as exc:  # pragma: no cover — the dep is pinned
        # Both routers map a message containing "not installed" to a 500, because
        # a missing server-side dependency is not the uploader's mistake.
        raise SpreadsheetError("Excel support is not installed on the server.") from exc

    # Authoritative at the point of use, not only at boot: a script or worker that
    # never ran the lifespan still cannot parse an undefused workbook.
    assert_xlsx_parser_hardened()

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises many types on a bad file
        raise SpreadsheetError(_UNREADABLE) from exc
    try:
        ws = wb.active
        if ws is None:
            return []
        # Bounded AS IT STREAMS. openpyxl's read-only reader does not buffer the
        # sheet, but the old loop appended every row it yielded into a list, so
        # the unbounded part was ours.
        return _bounded(
            ["" if c is None else str(c) for c in row]
            for row in ws.iter_rows(values_only=True)
        )
    except MemoryError:
        # NOT the uploader's mistake. MemoryError subclasses Exception, so the
        # clause below would have filed a resource-exhaustion event as "is it a
        # valid .xlsx?" and told nobody (security review 2026-10-06).
        raise
    except Exception as exc:  # noqa: BLE001 — openpyxl raises many types mid-stream
        # The load above carries this same handler; the streaming half never did,
        # so anything raised here left as itself and the route answered 500 on a
        # file the uploader could have fixed (review 2026-10-06).
        #
        # What actually reaches here: a workbook that LOADS and then fails while
        # streaming — a dangling shared-string reference is the realistic one
        # (openpyxl's `_reader.py` indexes `self.shared_strings[int(value)]`). An
        # earlier version of this comment also credited an entity-laced part, and
        # that was wrong: measured, a laced sheet fails inside `load_workbook` and
        # is caught by the handler above, not this one.
        #
        # WARNING, not info: neither caller logs — both do
        # `raise HTTPException(...) from exc` and the chained traceback is
        # discarded — so this line is the only trace, and a repeating openpyxl
        # regression should cross a default alert threshold rather than look like
        # a user's bad file.
        log.warning("question_import.workbook_unreadable", error_type=type(exc).__name__)
        raise SpreadsheetError(_UNREADABLE) from exc
    finally:
        wb.close()


def read_spreadsheet(filename: str, content: bytes) -> list[list[str]]:
    """Read an uploaded .xlsx or .csv into string rows. Raises SpreadsheetError.

    Every cell comes back as a string, including numbers — a "Points" column
    Excel stored as 2.0 must not become the string "2.0" to the caller's int
    parser, so the row parsers below go through float() for exactly that.

    DISPATCH IS BY CONTENT, NOT BY ``filename``. This used to branch on
    ``filename.endswith(".csv")`` with openpyxl as the fall-through, and
    ``filename`` comes straight off the multipart part — so the uploader chose
    which parser ran (review 2026-10-06). See :func:`looks_like_xlsx` for why the
    predicate asks "is this a workbook?" rather than "is this text?", and why that
    cannot refuse a file which imports today.

    ``filename`` still has one job, and it is not dispatch: it picks the SENTENCE
    a failure gets. A file called ``.xlsx`` whose bytes are not a workbook is a
    wrong or corrupt upload, and naming the two formats that work is more use than
    reading it as a one-column CSV and reporting "no question rows found" two
    layers up.
    """
    if looks_like_xlsx(content):
        return _read_workbook(content)
    if _claims_to_be_a_workbook(filename):
        # The bytes already decided; this only chooses the wording.
        raise SpreadsheetError(_UNREADABLE)
    return _read_csv(content)


def _bounded(source: Iterable[list[str]]) -> list[list[str]]:
    """Collect rows, slicing each to :data:`MAX_COLS` and stopping early.

    Stops at ``MAX_ROWS + 1`` rows with content — one past the cap, so the
    parser can still report "stopped after N rows" rather than silently
    truncating — or at :data:`MAX_SCAN_ROWS` rows examined, whichever comes
    first. Blank rows are kept in place so the row numbers the parser reports
    still match the spreadsheet a person is looking at.
    """
    out: list[list[str]] = []
    used = 0
    for scanned, row in enumerate(source, start=1):
        trimmed = row[:MAX_COLS]
        out.append(trimmed)
        if any(c.strip() for c in trimmed):
            used += 1
            if used > MAX_ROWS:
                break
        if scanned >= MAX_SCAN_ROWS:
            break
    return out


def _skip_header(rows: list[list[str]]) -> int:
    """1 when row 1 is the template header, else 0.

    This USED to return 1 whenever the first cell merely *contained*
    "question", which silently ate a real first row whose prompt had the word
    in it — "Which question type is this?", or any sheet starting with a
    question about questions. Nothing reported it: the row did not become an
    error, it became the header.

    So: an exact "question" (what the template writes), or a labelled variant
    like "Question Text" ONLY when the next cell also looks like a column
    heading. A genuine question's second cell holds an option's text, so it
    takes both halves being wrong at once to lose a row now.
    """
    if not rows or not rows[0]:
        return 0
    first = (rows[0][0] or "").strip().casefold()
    second = (rows[0][1] or "").strip().casefold() if len(rows[0]) > 1 else ""
    if first == "question":
        return 1
    if first.startswith("question") and second.startswith("option"):
        return 1
    return 0


def _points(raw: str) -> int:
    """1-100. An unreadable value is 1 rather than an error: a wrong weight is
    worth flagging, but it is not worth rejecting the question text over."""
    if not raw:
        return 1
    try:
        return max(1, min(100, int(float(raw))))
    except ValueError:
        return 1


def parse_mcq_rows(rows: list[list[str]]) -> tuple[list[dict[str, Any]], list[RowError]]:
    """The seven-column exam layout. Returns (items, per-row errors).

    An item is ``{row, prompt, options, correct_index, points}``. A blank line
    is skipped silently; anything else that cannot be used comes back as a
    RowError naming the 1-based row, because "row 14: correct answer 'E' matches
    no option" is actionable and "import failed" is not.

    ``row`` travels WITH the item, not just with the errors. A caller that
    validates further (an exam builds a Pydantic model from each item) can then
    still name the line it rejected; deriving it afterwards means matching on
    prompt text, which breaks on the duplicate prompts a real sheet contains.
    """
    items: list[dict[str, Any]] = []
    errors: list[RowError] = []
    start = _skip_header(rows)
    # Counts only rows with CONTENT against the cap. Excel emits a trailing
    # blank row for every cell anyone ever touched, so counting blanks let a
    # sheet with 1,500 empty rows push a real question past the limit and drop
    # it behind a generic "stopped after 1000 rows" (review 2026-10-05).
    used = 0
    # Bounded here TOO, and not only by `used`: the reader is the real bound
    # (see _bounded), but a caller handing this function an unbounded list
    # directly must not be able to spin it. Stripping 16,384 cells per row for a
    # million rows is the shape that got measured at 12.5 s.
    for i in range(start, min(len(rows), start + MAX_SCAN_ROWS)):
        rownum = i + 1  # 1-based, as the spreadsheet shows it
        cells = [(c or "").strip() for c in rows[i]]
        if not any(cells):
            continue
        used += 1
        if used > MAX_ROWS:
            errors.append(RowError(
                rownum,
                f"stopped after {MAX_ROWS} rows — split the file and import the rest",
            ))
            break
        prompt = cells[0] if cells else ""
        # Both lists, deliberately. `columns` keeps the four Option cells in
        # place so a letter means a column; `options` is the blanks-removed list
        # that gets stored. Conflating the two is the defect in
        # correct_to_index's docstring.
        columns = [cells[j] if j < len(cells) else "" for j in range(1, 5)]
        options = [c for c in columns if c]
        correct_raw = cells[5] if len(cells) > 5 else ""
        if not prompt:
            errors.append(RowError(rownum, "missing question text"))
            continue
        if any("\x00" in c for c in cells):
            # Postgres refuses a NUL in a text column, and it would surface as a
            # 500 at flush time — after the other rows had been built. Named
            # here instead, like every other unusable row.
            errors.append(RowError(rownum, "contains a NUL byte"))
            continue
        if len(options) < 2:
            errors.append(RowError(rownum, "need at least 2 options"))
            continue
        try:
            correct_index = correct_to_index(correct_raw, options, columns=columns)
        except ValueError as exc:
            errors.append(RowError(rownum, str(exc)))
            continue
        items.append({
            "row": rownum,
            "prompt": prompt, "options": options, "correct_index": correct_index,
            "points": _points(cells[6] if len(cells) > 6 else ""),
        })
    return items, errors


def parse_bank_rows(
    rows: list[list[str]],
    *,
    default_difficulty: str = "medium",
    default_language: str = "en",
) -> tuple[list[dict[str, Any]], list[RowError]]:
    """The bank layout: the exam's seven columns plus optional Difficulty,
    Language and Competency.

    Returns items shaped for ``question_banks.create_question`` — with ``kind``
    fixed to ``mcq``, because a coding question needs test cases that no
    spreadsheet row can express. The defaults apply per row, so a sheet that
    fills Difficulty for some questions and leaves it blank for others gets the
    caller's choice only where the file is silent.

    An unrecognised Difficulty or Language is an ERROR, not a silent fallback.
    Both are graded facts: "hard" questions carry the weight of hard questions
    in a picker, and a Telugu question shown to an English candidate is an
    unusable question. Quietly importing "Hrad" as medium/en would put a wrong
    fact in a reusable library, which is worse than a row the importer is told
    to fix.
    """
    base, errors = parse_mcq_rows(rows)
    if not base:
        return [], errors

    items: list[dict[str, Any]] = []
    for item in base:
        rownum = int(item["row"])
        cells = [(c or "").strip() for c in rows[rownum - 1]]
        difficulty = (cells[7] if len(cells) > 7 else "").strip().casefold() or default_difficulty
        language = (cells[8] if len(cells) > 8 else "").strip().casefold() or default_language
        competency = (cells[9] if len(cells) > 9 else "").strip()
        if difficulty not in DIFFICULTIES:
            errors.append(RowError(
                rownum, f"difficulty '{difficulty}' is not easy, medium or hard"
            ))
            continue
        if language not in LANGUAGES:
            errors.append(RowError(rownum, f"language '{language}' is not en, hi or te"))
            continue
        out: dict[str, Any] = {
            **item, "kind": "mcq", "difficulty": difficulty, "language": language,
        }
        if competency:
            out["competencies"] = [
                {"id": competency_slug(competency), "name": competency}
            ]
        items.append(out)
    errors.sort(key=lambda e: e.row)
    return items, errors
